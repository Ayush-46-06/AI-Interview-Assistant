import os
from pathlib import Path
from fastapi import HTTPException, status
import pypdf
import docx

MAX_RESUME_TEXT_LENGTH = 50000

def validate_magic_bytes(file_path: Path, extension: str) -> None:
    """Validate file signature (magic bytes) to prevent extension spoofing."""
    try:
        with open(file_path, "rb") as f:
            header = f.read(5)
    except Exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read file signature.")
        
    if extension == ".pdf":
        if not header.startswith(b"%PDF-"):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid PDF file signature.")
    elif extension == ".docx":
        if not header.startswith(b"PK\x03\x04"):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid DOCX file signature.")

def parse_pdf(file_path: Path) -> str:
    text_blocks = []
    try:
        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_blocks.append(extracted.strip())
    except Exception as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Failed to parse PDF: {str(e)}")
        
    text = "\n".join(text_blocks).strip()
    if not text:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, 
            detail="Could not extract text from PDF. The file may be image-only/scanned. OCR is deferred."
        )
    return text

def parse_docx(file_path: Path) -> str:
    text_blocks = []
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                text_blocks.append(para.text.strip())
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_blocks.append(" | ".join(row_text))
    except Exception as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Failed to parse DOCX: {str(e)}")
        
    text = "\n".join(text_blocks).strip()
    if not text:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, 
            detail="Could not extract text from DOCX. File may be empty."
        )
    return text

def parse_resume_file(file_path: Path, filename: str) -> str:
    """
    Parses a locally saved PDF or DOCX file, validates its signature,
    and returns the extracted text safely.
    """
    ext = os.path.splitext(filename)[1].lower()
    
    # 1. Magic byte validation
    validate_magic_bytes(file_path, ext)
    
    # 2. Text Extraction
    if ext == ".pdf":
        extracted_text = parse_pdf(file_path)
    elif ext == ".docx":
        extracted_text = parse_docx(file_path)
    else:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported file extension inside parser.")
        
    # 3. Limit Enforcement
    if len(extracted_text) > MAX_RESUME_TEXT_LENGTH:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, 
            detail=f"Extracted resume text ({len(extracted_text)} chars) exceeds the {MAX_RESUME_TEXT_LENGTH} character limit."
        )
        
    return extracted_text
